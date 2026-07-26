import os
import json
import json_repair
import base64
import httpx
from typing import Dict, Any, List
import csv
from io import StringIO
import openpyxl
import datetime
import time
import asyncio
import logging

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

GEMINI_API_KEY = os.environ.get("GEMINI_API_KEY") or os.environ.get("GOOGLE_API_KEY")
MODEL_NAME = "gemini-3.6-flash"
GEMINI_MAX_CONCURRENT = int(os.environ.get("GEMINI_MAX_CONCURRENT", "3"))

PROMPT_V1 = """
You are an expert courier tariff data extractor.
Your job is to read a semantic JSON representation of a document and extract the tariff data.
The JSON provides 'sheet_metadata' (including 'merged_regions') and specific extracted 'rows'.

BUSINESS CONTEXT:
This extracted data will power a courier quoting engine. The quoting engine needs to calculate and compare shipping prices across different vendors (like FedEx, DHL). 
It is CRITICAL that Terms and Conditions (T&Cs) are correctly categorized because some rules apply to the entire document, some apply only to a specific zone, and some apply only to a specific country. If a zone-specific or country-specific rule is extracted as a document-wide rule, the quoting engine will show incorrect rules to the customer!

CRITICAL EXTRACTION STRATEGY:
Step 1: Identify Document Structure
- Scan the blocks to locate Titles, the main Data Tables, and Footer Notes.
- Use `merged_regions` to understand if a header spans multiple columns (e.g. if C2:D2 is merged, the header in C2 applies to columns C and D).

Step 2: Parse Tables
- Find the rate matrices. Identify the Weight column and the Destination Zone headers.
- HORIZONTAL MATRICES: If you encounter a horizontal matrix where Countries/Zones are listed down the rows, and Weights are listed across the columns (e.g., 6, 8, 11+ as headers), you MUST transpose this data mentally. Our schema strictly requires `weight` to be the numeric weight and `zone` to be the country/region string. DO NOT output a zone name as a weight. Extract the data by mapping the column weight header to the row zone name.
- CONDITIONAL ZONES / MODIFIERS: If you encounter columns for the same zone but with different conditions (e.g., "Australia (With Food)" vs "Australia (Without Food)"), DO NOT merge this condition into the zone name. Instead, extract them as completely separate `sections` with different `service` names (e.g., Section 1: Service="With Food", Zone="Australia"; Section 2: Service="Without Food", Zone="Australia").
- MULTIPLE CARRIERS: If a single rate table lists different carriers (e.g. DPD, ARAMEX, FedEx) as the column headers instead of zones, you MUST extract them as completely separate `sections`. For example, if the table heading is "UK RATE LIST" and columns are "DPD" and "ARAMEX", create one section with `carrier`="DPD", `zone`="UK", and another section with `carrier`="ARAMEX", `zone`="UK". Do NOT treat carrier names as zones!
- TWO-TABLE FLAT+PER_KG PATTERN: Some sheets contain two stacked rate tables. The first table lists FLAT rates for weights up to a threshold (e.g., 0.5kg–30kg). The second table, introduced by a heading like "Multiplier rate per 1 KG from X KG" or "Per KG rate", lists PER_KG rates for weight brackets above that threshold (e.g., "30.1-50", "50.1-100"). When you see this pattern, extract BOTH tables as rates within the SAME section. For the first table, set price_type="FLAT" for those weight zones. For the second table rows with bracket weights (e.g. "30.1-50"), set price_type="PER_KG" and extract the weight EXACTLY as the bracket string (e.g. "30.1-50"). The heading row itself (e.g. "Multiplier rate per 1 KG from 30.1 KG") is your signal — use it as the evidence for your zone_segments PER_KG classification.

Step 3: Parse Notes & Rules
- Locate terms and conditions (T&Cs).
- If T&Cs appear directly below a specific column or are merged under a specific destination/zone, associate them ONLY with that destination/zone by setting scope='zone' and applies_to=[zone_name].
- If T&Cs are found at the bottom of a specific rate sheet (e.g. the "RATES" tab), they belong to the 'sheet' scope. These will only apply to services extracted from this exact sheet.
- If T&Cs are found on a dedicated "Index", "Cover", or "T&C" sheet that explicitly declares they apply to the entire file, use the 'document' scope.
- If T&Cs apply to a specific matrix on a sheet, use the 'section' scope.
- If you find text defining which countries belong to a specific zone, put it in `notes` with `scope="zone"` and `applies_to=["Zone Name"]`. YOU MUST FORMAT the `text` of this note exactly as a pipe-separated list of countries (e.g., "Belgium|Denmark|France"). Do not include prefixes like "Countries:" or the zone name in the text itself.
- CATEGORIZE every note into one of four priority levels:
  1. "CRITICAL": Strict restrictions, prohibited items, dangerous goods, seizure warnings, or rules about medicine/ghee.
  2. "BILLING": Surcharges, RTO penalties, duty taxes, or financial terms.
  3. "POINTER": If a note explicitly states that rules from another sheet apply here (e.g., "Refer to RATES sheet for terms"). Put the target sheet name in `applies_to`.
  4. "INFO": Standard delivery terms, transit times, or general information.

Step 4: Parse Zone Mappings (CRITICAL HYBRID ARCHITECTURE)
You will receive a compressed view of the sheet. Massive postcode mappings (thousands of rows) have been truncated.
- If you find a zone mapping table (e.g. mapping postcodes or countries to zones), you must declare it in `zone_mappings`.
- CRITICAL - DISTINGUISH TABLE TYPE:
  - If the key column contains COUNTRY NAMES (e.g. "Australia", "United Kingdom"), the column role MUST be `"country"`, NOT `"postcode"`. Use `"country"` for any table where countries map to zones.
  - If the key column contains NUMERIC POSTCODES or suburb names, use `"postcode"` or `"suburb"` accordingly.
- Mode 1 (AI): If it's a small list (under ~100 rows) and you can see it in the compressed view, set `mode: "AI"` and populate the `mapping` array with objects like `{"key_type": "country", "key": "Australia", "zone": "3"}` for countries or `{"key_type": "postcode", "key": "3000", "zone": "Metro"}` for postcodes.
- Mode 2 (PYTHON): If it's a massive postcode list or runs off the compressed view, DO NOT extract the rows. Set `mode: "PYTHON"` and return the Table Definition: `bounds` (only `top`, `left`, `right` needed) and `column_roles` (map the column letters like "A": "country", "B": "zone" OR "A": "postcode", "B": "suburb", "C": "zone"). If there is a column for remarks, no-service flags, or notes, label it "note". Python will use this definition to extract the raw data deterministically! If the table is a dense grid with repeating column pairs (e.g., Postcode in Col M, Zone in Col N AND Postcode in Col O, Zone in Col P), include ALL of those letters in `column_roles` and set `left` to the first column and `right` to the last column.

Step 5: Output JSON
- CRITICAL: Extract EVERY SINGLE RATE ROW from the table exactly as it appears. DO NOT skip any rows. DO NOT summarize or truncate patterns (e.g. if you see weights 6, 7, 8, 9, you MUST extract every single one). Skipping rows will cause catastrophic quoting errors.
- WEIGHT BRACKETS: If a weight cell contains a range like "30.1-50" or "100.1-300", extract it EXACTLY as that string. Do NOT expand it. Our backend handles expansion.
- Determine pricing models (FLAT vs PER_KG). Because pricing models can change per zone and per weight bracket (e.g. 1-10kg is FLAT, 11+ is PER_KG), extract a `zone_segments` array for each section. Determine this from explicit evidence in the document (headers, table titles, notes such as "Per Kg", "Rate/Kg", "Additional Kg", or a heading like "Multiplier rate per 1 KG from X KG"). If there is no clear evidence, return "UNKNOWN" rather than guessing.
- Return the EXACT schema below.

JSON SCHEMA:
{
  "sections": [
    {
      "carrier": "Carrier name (e.g. FedEx, DHL)",
      "service": "Service name. If missing from headers, attempt to infer it from the `sheet_name` (e.g., 'UK DPD Rate' -> 'DPD Rate'). If neither contains a valid service name (e.g., 'Sheet1'), leave it empty (\"\"). Do not guess.",
      "valid_from": "YYYY-MM-DD (if found, else null)",
      "valid_to": "YYYY-MM-DD (if found, else null)",
      "zone_segments": [
        {
          "zone": "Zone name (e.g. Zone 1, USA)",
          "transit_days": "String: Delivery period/transit time if explicitly provided for this zone, else null (e.g. '3-5', '3')",
          "segments": [
            {
              "start_weight": numeric (e.g. 1),
              "end_weight": numeric or null (e.g. 10 or null),
              "price_type": "FLAT|PER_KG|UNKNOWN",
              "confidence": numeric (0.0 to 1.0),
              "reason": "AI explanation for this segment"
            }
          ]
        }
      ],
      "rates": [
        {
          "weight": "Numeric weight OR string bracket (e.g. '21-30')",
          "zone": "Zone identifier",
          "price": "Numeric price (e.g. 15.50)"
        }
      ],
      "notes": [
        {
          "text": "Note text",
          "category": "CRITICAL|BILLING|POINTER|INFO",
          "scope": "document|sheet|section|zone",
          "applies_to": ["Zone Name"]
        }
      ],
      "zone_mappings": [
        {
          "mode": "AI|PYTHON",
          "confidence": numeric (0.0 to 1.0),
          "mapping": [{"key_type": "country|postcode|suburb", "key": "Australia or 3000", "zone": "3"}],
          "bounds": {"top": numeric, "left": "ColLetter", "right": "ColLetter"},
          "column_roles": {"A": "country", "B": "zone"}
        }
      ]
    }
  ]
}

Ensure your output is just raw JSON, with no markdown formatting or backticks around it.
"""

def get_excel_extraction_units(file_path: str, allowed_sheets: List[str] = None, skip_middle_sheets: List[str] = None, force_all_sheets: List[str] = None) -> List[Dict[str, Any]]:
    """Reads an Excel file and converts it into a list of extraction units (one per sheet) with compression."""
    wb = openpyxl.load_workbook(file_path, data_only=True)
    units = []
    
    skip_middle_sheets = skip_middle_sheets or []
    force_all_sheets = force_all_sheets or []
    violations = {}
    
    for sheet_name in wb.sheetnames:
        if allowed_sheets is not None and sheet_name not in allowed_sheets:
            continue
        sheet = wb[sheet_name]
        
        merged_row_indices = set()
        merged_regions_str = []
        for r in sheet.merged_cells.ranges:
            merged_regions_str.append(str(r))
            for row_idx in range(r.min_row, r.max_row + 1):
                merged_row_indices.add(row_idx)
                
        max_r = sheet.max_row or 1
        max_c = min(sheet.max_column, 500) if sheet.max_column else 500
        
        # Determine hidden rows and columns
        hidden_rows = set()
        for r_idx in range(1, max_r + 1):
            if r_idx in sheet.row_dimensions and sheet.row_dimensions[r_idx].hidden:
                hidden_rows.add(r_idx)
                
        hidden_cols = set()
        for _, dim in sheet.column_dimensions.items():
            if getattr(dim, 'hidden', False):
                min_c = getattr(dim, 'min', None)
                dim_max_c = getattr(dim, 'max', None)
                if min_c and dim_max_c:
                    for c_idx in range(min_c, dim_max_c + 1):
                        hidden_cols.add(c_idx)

        # Parse all visible rows to find the true bottom and extract data
        visible_rows_data = []
        for row_idx in range(1, max_r + 1):
            if row_idx in hidden_rows:
                continue
                
            row_dict = {}
            has_data = False
            for col_idx in range(1, max_c + 1):
                if col_idx in hidden_cols:
                    continue
                    
                cell = sheet.cell(row=row_idx, column=col_idx)
                val = cell.value
                if val is not None and str(val).strip() != "":
                    has_data = True
                    if isinstance(val, (datetime.datetime, datetime.date)):
                        if val.day == 1:
                            val = f"{val.month}-{str(val.year)[-2:]}"
                        else:
                            val = f"{val.day}-{val.month}"
                    elif isinstance(val, float):
                        if val.is_integer():
                            val = int(val)
                        else:
                            val = round(val, 1)
                        
                    col_letter = openpyxl.utils.get_column_letter(col_idx)
                    row_dict[col_letter] = str(val).strip()
                    
            if has_data:
                visible_rows_data.append({"row": row_idx, "data": row_dict})
                
        if not visible_rows_data:
            continue
            
        # Compression logic: top 150, bottom 150, and middle anomalies
        total_visible = len(visible_rows_data)
        if total_visible <= 300:
            compressed_data = visible_rows_data
        else:
            top_150 = visible_rows_data[:150]
            bottom_150 = visible_rows_data[-150:]
            
            if sheet_name in skip_middle_sheets:
                compressed_data = top_150 + bottom_150
                compressed_data.sort(key=lambda x: x["row"])
            else:
                middle = visible_rows_data[150:-150]
                anomalies = []
                for row_obj in middle:
                    r_idx = row_obj["row"]
                    data = row_obj["data"]
                    
                    # Anomaly 1: Merged cell
                    if r_idx in merged_row_indices:
                        anomalies.append(row_obj)
                        continue
                        
                    # Anomaly 2: Sparsity (1 or 2 filled cells, likely a rule/title)
                    if len(data) <= 2:
                        anomalies.append(row_obj)
                        continue
                        
                    # Anomaly 3: Long string (potential T&C)
                    has_long_string = any(isinstance(v, str) and len(v) > 40 for v in data.values())
                    if has_long_string:
                        anomalies.append(row_obj)
                        continue
                        
                compressed_data = top_150 + anomalies + bottom_150
                # Sort just in case anomalies overlapped with top/bottom (though they shouldn't by slice logic)
                compressed_data.sort(key=lambda x: x["row"])
                
            # Deduplicate just in case
            seen = set()
            deduped = []
            for item in compressed_data:
                if item["row"] not in seen:
                    seen.add(item["row"])
                    deduped.append(item)
            compressed_data = deduped
            
            if len(compressed_data) > 500:
                if sheet_name not in skip_middle_sheets and sheet_name not in force_all_sheets:
                    violations[sheet_name] = len(compressed_data)
                
            
        sheet_metadata = {
            "sheet_name": sheet_name,
            "merged_regions": merged_regions_str,
            "hidden_rows_skipped": list(hidden_rows)[:50], # Cap to avoid huge arrays
            "hidden_cols_skipped": [openpyxl.utils.get_column_letter(c) for c in hidden_cols],
            "total_visible_rows": total_visible,
            "compressed_rows_sent": len(compressed_data)
        }
        
        sheet_data = {
            "sheet_metadata": sheet_metadata,
            "rows": compressed_data
        }
            
        units.append({
            "source_type": "sheet",
            "source_name": sheet_name,
            "data": sheet_data
        })
        
    if violations:
        raise ValueError(f"TOO_MANY_ROWS|{json.dumps(violations)}")
        
    return units

async def call_gemini_api_with_retries(parts: List[Dict], unit_name: str, model: str = None) -> Dict[str, Any]:
    active_model = model if model else MODEL_NAME
    url = f"https://generativelanguage.googleapis.com/v1beta/models/{active_model}:generateContent?key={GEMINI_API_KEY}"
    body = {
        "contents": [{"role": "user", "parts": parts}],
        "generationConfig": {
            "temperature": 0.1,
            "responseMimeType": "application/json"
        }
    }
    
    max_retries = 3
    retry_count = 0
    start_time = time.time()
    
    async with httpx.AsyncClient(timeout=300.0) as client:
        while retry_count <= max_retries:
            try:
                response = await client.post(url, json=body)
                
                # Retry on rate limits (429) and server errors (500+)
                if response.status_code == 429 or response.status_code >= 500:
                    retry_count += 1
                    if retry_count > max_retries:
                        raise RuntimeError(f"Gemini API Error after retries: {response.status_code} {response.text}")
                    await asyncio.sleep(2 ** retry_count) # Exponential backoff
                    continue
                    
                # Do NOT retry on 400 bad requests
                if response.status_code != 200:
                    raise RuntimeError(f"Gemini API Error: {response.status_code} {response.text}")
                    
                data = response.json()
                
                # Parse usage metrics
                usage = data.get("usageMetadata", {})
                prompt_tokens = usage.get("promptTokenCount", 0)
                resp_tokens = usage.get("candidatesTokenCount", 0)
                
                text_content = ""
                for candidate in data.get("candidates", []):
                    for part in candidate.get("content", {}).get("parts", []):
                        text_content += part.get("text", "")
                        
                # Clean markdown
                text_content = text_content.strip()
                if text_content.startswith("```json"): text_content = text_content[7:]
                if text_content.startswith("```"): text_content = text_content[3:]
                if text_content.endswith("```"): text_content = text_content[:-3]
                text_content = text_content.strip()
                
                raw_json = {}
                try:
                    raw_json = json_repair.loads(text_content)
                    if not isinstance(raw_json, dict):
                        # In case json_repair returns a string or list when a dict was expected
                        import re
                        match = re.search(r'\{.*\}', text_content, re.DOTALL)
                        if match:
                            raw_json = json_repair.loads(match.group(0))
                            
                    # Normalize weight brackets: store as weight_min/weight_max instead of expanding
                    if isinstance(raw_json, dict) and "sections" in raw_json:
                        for s in raw_json["sections"]:
                            if "rates" in s:
                                normalized_rates = []
                                for r in s["rates"]:
                                    w = r.get("weight")
                                    w_str = str(w).strip() if w is not None else ""
                                    
                                    # Strip units
                                    w_clean = w_str.replace("kg", "").replace("KG", "").strip()
                                    
                                    import re as _re
                                    
                                    if isinstance(w, str) and "-" in w_clean:
                                        # Bracket: "30.1-50", "0.5-30"
                                        nums = _re.findall(r'\d+(?:\.\d+)?', w_clean)
                                        if len(nums) == 2:
                                            r["weight"] = float(nums[0])       # weight_min
                                            r["weight_max"] = float(nums[1])   # weight_max
                                            normalized_rates.append(r)
                                        # else drop
                                    elif isinstance(w, str) and "+" in w_clean:
                                        # Open-ended: "30+", "30.1+"
                                        nums = _re.findall(r'\d+(?:\.\d+)?', w_clean)
                                        if nums:
                                            r["weight"] = float(nums[0])
                                            r["weight_max"] = 99999.0
                                            normalized_rates.append(r)
                                    else:
                                        # Plain number: 5, 0.5, 30
                                        try:
                                            r["weight"] = float(w_clean)
                                            r["weight_max"] = None  # point rate
                                            normalized_rates.append(r)
                                        except (ValueError, TypeError):
                                            pass
                                
                                # Deduplicate by (zone, weight_min, weight_max)
                                unique_rates = {}
                                for r in normalized_rates:
                                    w_min = r.get("weight")
                                    w_max = r.get("weight_max")
                                    z = r.get("zone")
                                    if w_min is not None and z is not None:
                                        key = f"{z}_{w_min}_{w_max}"
                                        unique_rates[key] = r
                                s["rates"] = list(unique_rates.values())
                                
                except Exception as e:
                    logger.error(f"Failed to parse JSON using json_repair: {e}")
                    raise RuntimeError(f"Failed to parse JSON from AI response: {e}")
                        
                latency = time.time() - start_time
                logger.info(f"[Metrics - {unit_name}] Prompt Tokens: {prompt_tokens}, Resp Tokens: {resp_tokens}, Latency: {latency:.2f}s, Retries: {retry_count}, Status: SUCCESS")
                return raw_json
                
            except (httpx.RequestError, httpx.TimeoutException) as e:
                # Retry on transient network errors
                retry_count += 1
                if retry_count > max_retries:
                    logger.error(f"[Metrics - {unit_name}] Status: FAILED after {retry_count} retries. Error: {str(e)}")
                    raise RuntimeError(f"Network error calling Gemini: {str(e)}")
                await asyncio.sleep(2 ** retry_count)
                
    return {}

async def extract_rates_from_document(file_path: str, filename: str, allowed_sheets: List[str] = None, skip_middle_sheets: List[str] = None, force_all_sheets: List[str] = None, ai_context: str = None, ai_model: str = None) -> Dict[str, Any]:
    if not GEMINI_API_KEY:
        raise RuntimeError("GEMINI_API_KEY is not set.")

    start_time = time.time()
    ext = os.path.splitext(filename)[1].lower()
    
    master_json = {"sections": []}
    all_source_text = ""
    
    if ext in [".xlsx", ".xls"]:
        # Chunked extraction pipeline
        units = get_excel_extraction_units(file_path, allowed_sheets, skip_middle_sheets, force_all_sheets)
        semaphore = asyncio.Semaphore(GEMINI_MAX_CONCURRENT)
        
        async def process_unit(unit):
            async with semaphore:
                sheet_name = unit.get("source_name", "Unknown")
                unit_name = f"{filename} - {sheet_name}"
                json_str = json.dumps(unit["data"], ensure_ascii=False)
                # Cap the string length just in case a single sheet is maliciously huge
                if len(json_str) > 300000:
                    json_str = json_str[:300000] + "\n...[TRUNCATED]"
                
                parts = [
                    {"text": PROMPT_V1},
                    {"text": f"\n\nDocument Data (JSON where 'data' maps column letters to cell values):\n{json_str}"}
                ]
                
                if ai_context:
                    parts.append({"text": f"\n\nUSER INSTRUCTIONS / CUSTOM CONTEXT:\n{ai_context}\nPlease strictly follow the user instructions above if they clarify ambiguous data."})
                
                unit_res = await call_gemini_api_with_retries(parts, unit_name, ai_model)
                
                # Inject provenance (source) into each section
                for sec in unit_res.get("sections", []):
                    sec["source"] = {"source_type": unit["source_type"], "source_name": sheet_name}
                    for r in sec.get("rates", []):
                        r["source_ref"] = {"sheet": sheet_name}
                        r["confidence"] = 1.0
                    
                return unit_res, json_str

        tasks = [process_unit(u) for u in units]
        results = await asyncio.gather(*tasks)
        
        for res, source_str in results:
            master_json["sections"].extend(res.get("sections", []))
            all_source_text += source_str + "\n\n"
            
    else:
        # Single-pass extraction for PDFs/Images
        parts = [{"text": PROMPT_V1}]
        
        mime_type = "application/pdf"
        if ext in [".png"]: mime_type = "image/png"
        elif ext in [".jpg", ".jpeg"]: mime_type = "image/jpeg"
        
        with open(file_path, "rb") as f:
            b64_data = base64.b64encode(f.read()).decode("utf-8")
            
        parts.append({
            "inlineData": {
                "mimeType": mime_type,
                "data": b64_data
            }
        })
        
        if ai_context:
            parts.append({"text": f"\n\nUSER INSTRUCTIONS / CUSTOM CONTEXT:\n{ai_context}\nPlease strictly follow the user instructions above if they clarify ambiguous data."})
        
        unit_res = await call_gemini_api_with_retries(parts, "PDF/Image", ai_model)
        for sec in unit_res.get("sections", []):
            sec["source"] = {"source_type": "file", "source_name": filename}
            for r in sec.get("rates", []):
                r["source_ref"] = {"sheet": filename}
                r["confidence"] = 1.0
            master_json["sections"].append(sec)
            
    # Attach metadata for validation and UI
    master_json["metadata"] = {}
    master_json["metadata"]["extraction_time_sec"] = time.time() - start_time
    master_json["metadata"]["source_text"] = all_source_text
    master_json["metadata"]["extracted_at"] = datetime.datetime.now().isoformat()
        
    # Post-process to guarantee zone_segments exist for UI
    for section in master_json.get("sections", []):
        rates = section.get("rates", [])
        unique_zones = set(str(r.get("zone")) for r in rates if r.get("zone"))
        
        if "zone_segments" not in section or not isinstance(section["zone_segments"], list):
            section["zone_segments"] = []
            
        existing_zones = set(str(zs.get("zone")) for zs in section["zone_segments"] if isinstance(zs, dict) and zs.get("zone"))
        missing_zones = unique_zones - existing_zones
        
        for z in missing_zones:
            section["zone_segments"].append({
                "zone": z,
                "segments": [
                    {
                        "start_weight": 0,
                        "end_weight": None,
                        "price_type": "UNKNOWN",
                        "confidence": 0.0,
                        "reason": "No explicit evidence extracted by AI."
                    }
                ]
            })
            
    return master_json
