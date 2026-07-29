import openpyxl
import logging

logger = logging.getLogger(__name__)

def process_deterministic_zones(file_path: str, raw_json: dict):
    """
    Phase 3: Python Deterministic Parser.
    Intercepts zone_mappings with mode: "PYTHON" and extracts data deterministically.
    """
    if not raw_json or not raw_json.get("sections"):
        return raw_json
        
    # Check if we need to open the workbook (only if at least one PYTHON mode exists)
    needs_python = False
    for sec in raw_json["sections"]:
        for zm in sec.get("zone_mappings", []):
            if zm.get("mode") == "PYTHON":
                needs_python = True
                break
                
    if not needs_python:
        return raw_json
        
    logger.info(f"Hybrid parser triggered for {file_path}")
    
    try:
        # data_only=True gets values instead of formulas
        wb = openpyxl.load_workbook(file_path, data_only=True)
    except Exception as e:
        logger.error(f"Failed to load workbook for hybrid parsing: {e}")
        return raw_json
        
    try:
        for section in raw_json["sections"]:
            source = section.get("source", {})
            if source.get("source_type") != "sheet":
                continue
                
            sheet_name = source.get("source_name", "")
            if not sheet_name:
                continue
                
            # ai.py injects the exact sheet name now
            actual_sheet = None
            if sheet_name in wb.sheetnames:
                actual_sheet = sheet_name
                    
            if not actual_sheet:
                logger.warning(f"Could not find exact sheet '{sheet_name}' in workbook")
                continue
                
            sheet = wb[actual_sheet]
            
            zone_mappings = section.get("zone_mappings", [])
            for zm in zone_mappings:
                if zm.get("mode") != "PYTHON":
                    continue
                    
                bounds = zm.get("bounds", {})
                roles = zm.get("column_roles", {})
                
                top = bounds.get("top")
                left_col = bounds.get("left")
                right_col = bounds.get("right")
                
                if not top or not left_col or not right_col or not roles:
                    logger.warning(f"Invalid bounds/roles for PYTHON mode extraction in sheet {sheet_name}")
                    continue
                    
                try:
                    left_idx = openpyxl.utils.column_index_from_string(left_col)
                    right_idx = openpyxl.utils.column_index_from_string(right_col)
                except Exception as e:
                    logger.error(f"Failed to parse column bounds: {e}")
                    continue
                    
                extracted_mapping = []
                consecutive_empty_rows = 0
                max_row = sheet.max_row or 1048576
                
                min_col_idx = min(openpyxl.utils.column_index_from_string(c) for c in roles.keys())
                start_role = roles[openpyxl.utils.get_column_letter(min_col_idx)].lower()
                
                raw_rows = []
                
                for row_idx in range(top, max_row + 1):
                    has_valid_record = False
                    current_record = {}
                    is_empty_row = True
                    
                    # Unroll dense grids dynamically
                    for col_idx in range(left_idx, right_idx + 1):
                        col_letter = openpyxl.utils.get_column_letter(col_idx)
                        if col_letter not in roles:
                            continue
                            
                        cell = sheet.cell(row=row_idx, column=col_idx)
                        val = cell.value
                        role = roles[col_letter].lower()
                        
                        if val is not None and str(val).strip() != "":
                            is_empty_row = False
                            current_record[role] = str(val).strip()
                            if role != start_role and role != "note":
                                has_valid_record = True
                                
                    if is_empty_row:
                        consecutive_empty_rows += 1
                        if consecutive_empty_rows > 20:
                            break
                        continue
                        
                    consecutive_empty_rows = 0
                    
                    if has_valid_record and start_role in current_record:
                        shared_val = current_record[start_role]
                        shared_note = current_record.get("note")
                        
                        for r, v in current_record.items():
                            if r != start_role and r != "note":
                                rec = {start_role: shared_val, "zone": v}
                                if shared_note: rec["note"] = shared_note
                                raw_rows.append(rec)
                                
                pc_map = {}
                suburb_map = {}
                state_map = {}
                country_map = {}
                
                def parse_postcodes(val_str):
                    val_str = str(val_str).replace(' ', '')
                    parts = val_str.split(',')
                    result = []
                    for p in parts:
                        if '-' in p:
                            sub = p.split('-')
                            if len(sub) == 2 and sub[0].isdigit() and sub[1].isdigit():
                                try:
                                    s = int(sub[0])
                                    e = int(sub[1])
                                    for i in range(s, e + 1):
                                        result.append(str(i))
                                except:
                                    result.append(p)
                            else:
                                result.append(p)
                        else:
                            result.append(p)
                    return result
                
                for r in raw_rows:
                    z = str(r["zone"]).strip().upper()
                    if not z or z == "NONE":
                        continue
                        
                    n = r.pop("note", None)
                    for note_alias in ["notes", "remark", "remarks", "no_service", "service"]:
                        if note_alias in r:
                            val = r.pop(note_alias)
                            if not n:
                                n = val
                                
                    if r.get("postcode"):
                        pcs = parse_postcodes(r["postcode"])
                        for pc in pcs:
                            pc_map[pc] = {"zone": z, "note": n}
                            
                    if r.get("suburb") or r.get("city"):
                        sub = str(r.get("suburb") or r.get("city")).strip().upper()
                        suburb_map[sub] = {"zone": z, "note": n}
                        
                    if r.get("state"):
                        st = str(r["state"]).strip().upper()
                        state_map[st] = {"zone": z, "note": n}
                        
                    if r.get("country"):
                        ct = str(r["country"]).strip().upper()
                        country_map[ct] = {"zone": z, "note": n}
                        
                extracted_mapping = []
                
                # Group consecutive numeric postcodes
                sorted_pcs = []
                alpha_pcs = []
                int_to_str = {}
                for pc in pc_map.keys():
                    if pc.isdigit():
                        val = int(pc)
                        sorted_pcs.append(val)
                        # Keep the first string representation if there are duplicates (e.g. 0800 and 800)
                        if val not in int_to_str:
                            int_to_str[val] = pc
                    else:
                        alpha_pcs.append(pc)
                sorted_pcs.sort()
                
                grouped_pcs = []
                if sorted_pcs:
                    start_pc = sorted_pcs[0]
                    prev_pc = sorted_pcs[0]
                    
                    for i in range(1, len(sorted_pcs)):
                        curr_pc = sorted_pcs[i]
                        if curr_pc == prev_pc + 1 and pc_map[int_to_str[curr_pc]] == pc_map[int_to_str[prev_pc]]:
                            prev_pc = curr_pc
                        else:
                            if start_pc == prev_pc:
                                grouped_pcs.append(int_to_str[start_pc])
                            else:
                                grouped_pcs.append(f"{int_to_str[start_pc]}-{int_to_str[prev_pc]}")
                            start_pc = curr_pc
                            prev_pc = curr_pc
                            
                    if start_pc == prev_pc:
                        grouped_pcs.append(int_to_str[start_pc])
                    else:
                        grouped_pcs.append(f"{int_to_str[start_pc]}-{int_to_str[prev_pc]}")
                        
                for g in grouped_pcs + alpha_pcs:
                    key_for_lookup = g.split('-')[0] if '-' in g else g
                    if key_for_lookup.isdigit() and int(key_for_lookup) in int_to_str:
                        key_for_lookup = int_to_str[int(key_for_lookup)]
                        
                    data = pc_map[key_for_lookup]
                    rec = {"key_type": "postcode", "key": g, "zone": data["zone"]}
                    if data["note"]: rec["note"] = data["note"]
                    extracted_mapping.append(rec)
                    
                for sub, data in suburb_map.items():
                    rec = {"key_type": "suburb", "key": sub, "zone": data["zone"]}
                    if data["note"]: rec["note"] = data["note"]
                    extracted_mapping.append(rec)
                    
                for st, data in state_map.items():
                    rec = {"key_type": "state", "key": st, "zone": data["zone"]}
                    if data["note"]: rec["note"] = data["note"]
                    extracted_mapping.append(rec)
                    
                for ct, data in country_map.items():
                    rec = {"key_type": "country", "key": ct, "zone": data["zone"]}
                    if data["note"]: rec["note"] = data["note"]
                    extracted_mapping.append(rec)
                    
                # Update the JSON payload
                zm["mapping"] = extracted_mapping
                # Don't change mode to AI, so the UI still knows it was a PYTHON-extracted massive dataset.
                
    finally:
        wb.close()
        
    return raw_json
